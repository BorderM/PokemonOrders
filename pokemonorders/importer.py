from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .db import connect, create_request, init_db, upsert_person


def cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def as_int(value: Any, default: int = 1) -> int:
    try:
        return int(float(value))
    except Exception:
        return default


def maybe_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except Exception:
        return None


def import_workbook(path: str | Path) -> dict:
    init_db()
    wb = load_workbook(path, data_only=True, read_only=True)
    summary = {"people": 0, "requests": 0, "pricing_entries": 0, "shipping_rules": 0, "sheets": wb.sheetnames}

    if "People" in wb.sheetnames:
        ws = wb["People"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            email, name = (list(row) + [None, None])[:2]
            if email:
                upsert_person(cell_value(email), cell_value(name) or None)
                summary["people"] += 1

    if "Orders_Clean" in wb.sheetnames:
        ws = wb["Orders_Clean"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            timestamp, name, email, card, quantity, language, condition, variant, seller, notes = (list(row) + [None] * 10)[:10]
            if not card or not email:
                continue
            create_request({
                "name": cell_value(name),
                "email": cell_value(email),
                "card_name_code": cell_value(card),
                "quantity": as_int(quantity),
                "language": cell_value(language) or "English",
                "minimum_condition": cell_value(condition) or "Near Mint",
                "variant": cell_value(variant) or "Default",
                "seller": cell_value(seller),
                "notes": cell_value(notes),
                "source": "xlsx:Orders_Clean",
            })
            summary["requests"] += 1

    for sheet_name in ["Pricing", "Pricing Manual", "Order_Ledger"] + [s for s in wb.sheetnames if s.startswith("Archive_")]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            continue
        headers = [cell_value(c) for c in first_row]
        header_map = {h.strip().lower(): i for i, h in enumerate(headers) if h}
        if not any(h in header_map for h in ["requested by", "requester"]):
            continue
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            values = list(row)
            requester_idx = header_map.get("requested by", header_map.get("requester", 0))
            requested_by = values[requester_idx] if requester_idx < len(values) else ""
            card_idx = header_map.get("card", header_map.get("card/product"))
            card = values[card_idx] if card_idx is not None and card_idx < len(values) else ""
            if not requested_by or not card:
                continue
            qty_idx = header_map.get("qty", header_map.get("quantity"))
            cond_idx = header_map.get("condition", header_map.get("requested condition"))
            variant_idx = header_map.get("variant")
            price_idx = header_map.get("price", header_map.get("unit price"))
            total_idx = header_map.get("total price")
            shipping_idx = header_map.get("shipping units")
            notes_idx = header_map.get("notes")
            with connect() as conn:
                conn.execute(
                    """INSERT INTO pricing_entries(
                       requester_name, quantity, card_product, requested_condition, variant,
                       unit_price, total_price, shipping_units, notes, source_sheet, source_row
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        cell_value(requested_by),
                        as_int(values[qty_idx]) if qty_idx is not None and qty_idx < len(values) else None,
                        cell_value(card),
                        cell_value(values[cond_idx]) if cond_idx is not None and cond_idx < len(values) else None,
                        cell_value(values[variant_idx]) if variant_idx is not None and variant_idx < len(values) else None,
                        maybe_float(values[price_idx]) if price_idx is not None and price_idx < len(values) else None,
                        maybe_float(values[total_idx]) if total_idx is not None and total_idx < len(values) else None,
                        maybe_float(values[shipping_idx]) if shipping_idx is not None and shipping_idx < len(values) else None,
                        cell_value(values[notes_idx]) if notes_idx is not None and notes_idx < len(values) else None,
                        sheet_name,
                        idx,
                    ),
                )
                summary["pricing_entries"] += 1

    if "Shipping_Rules" in wb.sheetnames:
        ws = wb["Shipping_Rules"]
        with connect() as conn:
            for row in ws.iter_rows(min_row=2, values_only=True):
                rule, units, priority, min_value, max_value, extra_units = (list(row) + [None] * 6)[:6]
                if not rule:
                    continue
                conn.execute(
                    """INSERT INTO shipping_rules(rule, units, priority, min_unit_value, max_unit_value, extra_units)
                       VALUES (?, ?, ?, ?, ?, ?)""",
                    (
                        cell_value(rule),
                        float(units or 0),
                        int(float(priority or 99)),
                        maybe_float(min_value),
                        maybe_float(max_value),
                        float(extra_units or 0),
                    ),
                )
                summary["shipping_rules"] += 1

    return summary
